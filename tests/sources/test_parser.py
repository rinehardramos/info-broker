"""Tests for app/sources/parser.py — TDD suite written before implementation."""

import csv
import os
import tempfile

import pandas as pd
import pytest

from app.sources.parser import (
    estimate_tokens,
    parse_csv,
    parse_docx,
    parse_excel,
    parse_file,
    parse_pdf,
    parse_txt,
)


# ---------------------------------------------------------------------------
# estimate_tokens
# ---------------------------------------------------------------------------


def test_estimate_tokens_basic():
    # "hello world" has len 11 → 11 // 4 = 2
    result = estimate_tokens("hello world")
    assert result == 2


def test_estimate_tokens_empty():
    assert estimate_tokens("") == 0


def test_estimate_tokens_longer():
    text = "a" * 100
    assert estimate_tokens(text) == 25


# ---------------------------------------------------------------------------
# parse_txt
# ---------------------------------------------------------------------------


def test_parse_txt_empty_returns_empty_list():
    assert parse_txt("", "test.txt") == []


def test_parse_txt_whitespace_only_returns_empty_list():
    assert parse_txt("   \n\n   ", "test.txt") == []


def test_parse_txt_three_short_paragraphs_titles():
    content = "Para one.\n\nPara two.\n\nPara three."
    findings = parse_txt(content, "report.txt")
    assert len(findings) >= 1
    for finding in findings:
        assert finding["title"].startswith("report.txt - chunk ")


def test_parse_txt_content_has_contextual_prefix():
    content = "Hello world paragraph."
    findings = parse_txt(content, "notes.txt")
    assert len(findings) == 1
    assert findings[0]["content"].startswith("[File: notes.txt |")


def test_parse_txt_finding_fields():
    content = "Some content here."
    findings = parse_txt(content, "data.txt")
    assert len(findings) == 1
    f = findings[0]
    assert f["source"] == "file_upload"
    assert f["confidence"] == 85
    assert "title" in f
    assert "content" in f


def test_parse_txt_chunk_numbering():
    # Build content that will definitely exceed 2000 chars to force multiple chunks.
    paragraph = "word " * 120  # ~600 chars each
    content = "\n\n".join([paragraph] * 5)
    findings = parse_txt(content, "big.txt")
    assert len(findings) >= 2
    titles = [f["title"] for f in findings]
    assert "big.txt - chunk 1" in titles
    assert "big.txt - chunk 2" in titles


def test_parse_txt_chunk_content_includes_chunk_tag():
    paragraph = "word " * 120
    content = "\n\n".join([paragraph] * 5)
    findings = parse_txt(content, "big.txt")
    for i, f in enumerate(findings, start=1):
        assert f"[File: big.txt | Chunk {i}]" in f["content"]


def test_parse_txt_all_findings_have_file_upload_source():
    paragraph = "word " * 120
    content = "\n\n".join([paragraph] * 5)
    findings = parse_txt(content, "big.txt")
    assert all(f["source"] == "file_upload" for f in findings)


# ---------------------------------------------------------------------------
# parse_csv
# ---------------------------------------------------------------------------


def _write_csv(rows: list[dict], fieldnames: list[str] | None = None) -> str:
    """Write rows to a temp CSV file and return the path."""
    fh = tempfile.NamedTemporaryFile(
        mode="w", suffix=".csv", delete=False, newline=""
    )
    if fieldnames is None and rows:
        fieldnames = list(rows[0].keys())
    writer = csv.DictWriter(fh, fieldnames=fieldnames or [])
    writer.writeheader()
    writer.writerows(rows)
    fh.close()
    return fh.name


def test_parse_csv_schema_is_first_finding_with_confidence_95():
    path = _write_csv(
        [{"name": "Alice", "age": 30}, {"name": "Bob", "age": 25}]
    )
    try:
        findings = parse_csv(path, "people.csv")
        assert findings[0]["confidence"] == 95
        assert findings[0]["title"] == "people.csv - Schema"
    finally:
        os.unlink(path)


def test_parse_csv_schema_contains_column_names():
    path = _write_csv([{"city": "NY", "pop": 8000000}])
    try:
        findings = parse_csv(path, "cities.csv")
        schema_content = findings[0]["content"]
        assert "city" in schema_content
        assert "pop" in schema_content
    finally:
        os.unlink(path)


def test_parse_csv_schema_contains_row_count():
    rows = [{"x": i} for i in range(7)]
    path = _write_csv(rows)
    try:
        findings = parse_csv(path, "numbers.csv")
        assert "7" in findings[0]["content"]
    finally:
        os.unlink(path)


def test_parse_csv_schema_content_has_contextual_prefix():
    path = _write_csv([{"a": 1}])
    try:
        findings = parse_csv(path, "sample.csv")
        assert findings[0]["content"].startswith("[File: sample.csv |")
    finally:
        os.unlink(path)


def test_parse_csv_five_rows_returns_schema_plus_one_row_finding():
    rows = [{"n": i} for i in range(5)]
    path = _write_csv(rows)
    try:
        findings = parse_csv(path, "five.csv")
        # schema + 1 row group (rows 1-5)
        assert len(findings) == 2
    finally:
        os.unlink(path)


def test_parse_csv_21_rows_returns_schema_plus_two_row_findings():
    rows = [{"n": i} for i in range(21)]
    path = _write_csv(rows)
    try:
        findings = parse_csv(path, "big.csv")
        # schema + group(1-20) + group(21-21)
        assert len(findings) == 3
    finally:
        os.unlink(path)


def test_parse_csv_row_findings_have_correct_titles():
    rows = [{"n": i} for i in range(25)]
    path = _write_csv(rows)
    try:
        findings = parse_csv(path, "data.csv")
        titles = [f["title"] for f in findings]
        assert "data.csv - rows 1-20" in titles
        assert "data.csv - rows 21-25" in titles
    finally:
        os.unlink(path)


def test_parse_csv_row_findings_confidence_is_85():
    rows = [{"n": i} for i in range(3)]
    path = _write_csv(rows)
    try:
        findings = parse_csv(path, "small.csv")
        row_findings = findings[1:]  # skip schema
        assert all(f["confidence"] == 85 for f in row_findings)
    finally:
        os.unlink(path)


def test_parse_csv_all_findings_have_file_upload_source():
    rows = [{"a": i, "b": i * 2} for i in range(5)]
    path = _write_csv(rows)
    try:
        findings = parse_csv(path, "ab.csv")
        assert all(f["source"] == "file_upload" for f in findings)
    finally:
        os.unlink(path)


def test_parse_csv_empty_file_returns_only_schema():
    path = _write_csv([], fieldnames=["col1", "col2"])
    try:
        findings = parse_csv(path, "empty.csv")
        assert len(findings) == 1
        assert findings[0]["title"] == "empty.csv - Schema"
        assert "0" in findings[0]["content"]
    finally:
        os.unlink(path)


def test_parse_csv_row_finding_content_has_contextual_prefix():
    rows = [{"x": i} for i in range(3)]
    path = _write_csv(rows)
    try:
        findings = parse_csv(path, "xy.csv")
        row_findings = findings[1:]
        for f in row_findings:
            assert f["content"].startswith("[File: xy.csv |")
    finally:
        os.unlink(path)


# ---------------------------------------------------------------------------
# parse_excel
# ---------------------------------------------------------------------------


def _write_xlsx(sheets: dict) -> str:
    """Write an xlsx file with named sheets and return the temp path."""
    import pandas as pd

    fh = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    fh.close()
    with pd.ExcelWriter(fh.name, engine="openpyxl") as writer:
        for sheet_name, rows in sheets.items():
            df = pd.DataFrame(rows)
            df.to_excel(writer, sheet_name=sheet_name, index=False)
    return fh.name


def test_parse_excel_returns_findings_for_each_sheet():
    path = _write_xlsx({"People": [{"name": "Alice", "age": 30}], "Cars": [{"make": "Ford"}]})
    try:
        findings = parse_excel(path, "data.xlsx")
        titles = [f["title"] for f in findings]
        assert any("People" in t for t in titles)
        assert any("Cars" in t for t in titles)
    finally:
        os.unlink(path)


def test_parse_excel_schema_finding_has_sheet_name_in_title():
    path = _write_xlsx({"Employees": [{"name": "Bob", "dept": "Eng"}]})
    try:
        findings = parse_excel(path, "report.xlsx")
        schema_titles = [f["title"] for f in findings if "Schema" in f["title"]]
        assert any("Employees" in t for t in schema_titles)
        assert any("report.xlsx" in t for t in schema_titles)
    finally:
        os.unlink(path)


def test_parse_excel_row_finding_has_sheet_name_in_title():
    rows = [{"n": i} for i in range(5)]
    path = _write_xlsx({"Sheet1": rows})
    try:
        findings = parse_excel(path, "numbers.xlsx")
        row_titles = [f["title"] for f in findings if "rows" in f["title"]]
        assert any("Sheet1" in t for t in row_titles)
    finally:
        os.unlink(path)


def test_parse_excel_all_findings_have_file_upload_source():
    path = _write_xlsx({"Data": [{"a": 1, "b": 2}]})
    try:
        findings = parse_excel(path, "test.xlsx")
        assert all(f["source"] == "file_upload" for f in findings)
    finally:
        os.unlink(path)


def test_parse_excel_schema_title_format():
    path = _write_xlsx({"MySheet": [{"x": 1}]})
    try:
        findings = parse_excel(path, "file.xlsx")
        schema_titles = [f["title"] for f in findings if "Schema" in f["title"]]
        assert "file.xlsx - MySheet Schema" in schema_titles
    finally:
        os.unlink(path)


def test_parse_excel_row_title_format():
    rows = [{"n": i} for i in range(3)]
    path = _write_xlsx({"Tab": rows})
    try:
        findings = parse_excel(path, "file.xlsx")
        row_titles = [f["title"] for f in findings if "rows" in f["title"]]
        assert "file.xlsx - Tab rows 1-3" in row_titles
    finally:
        os.unlink(path)


# ---------------------------------------------------------------------------
# parse_pdf
# ---------------------------------------------------------------------------


def _write_pdf(pages: list) -> str:
    """Create a temp PDF with one page per string in pages."""
    from fpdf import FPDF

    pdf = FPDF()
    for text in pages:
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.multi_cell(0, 10, text)
    fh = tempfile.NamedTemporaryFile(suffix=".pdf", delete=False)
    fh.close()
    pdf.output(fh.name)
    return fh.name


def test_parse_pdf_one_finding_per_page():
    path = _write_pdf(["Page one content", "Page two content"])
    try:
        findings = parse_pdf(path, "doc.pdf")
        assert len(findings) == 2
    finally:
        os.unlink(path)


def test_parse_pdf_page_numbers_in_titles():
    path = _write_pdf(["First page", "Second page"])
    try:
        findings = parse_pdf(path, "report.pdf")
        titles = [f["title"] for f in findings]
        assert "report.pdf - Page 1" in titles
        assert "report.pdf - Page 2" in titles
    finally:
        os.unlink(path)


def test_parse_pdf_content_has_prefix():
    path = _write_pdf(["Hello world"])
    try:
        findings = parse_pdf(path, "notes.pdf")
        assert findings[0]["content"].startswith("[File: notes.pdf | Page 1]")
    finally:
        os.unlink(path)


def test_parse_pdf_all_findings_have_file_upload_source():
    path = _write_pdf(["Content here"])
    try:
        findings = parse_pdf(path, "test.pdf")
        assert all(f["source"] == "file_upload" for f in findings)
    finally:
        os.unlink(path)


def test_parse_pdf_empty_pages_skipped():
    path = _write_pdf(["Real content", "   ", "More content"])
    try:
        findings = parse_pdf(path, "mixed.pdf")
        # Page 2 is whitespace-only — should be skipped
        assert len(findings) == 2
    finally:
        os.unlink(path)


def test_parse_pdf_content_includes_page_text():
    path = _write_pdf(["Unique marker text 12345"])
    try:
        findings = parse_pdf(path, "marker.pdf")
        assert "Unique marker text 12345" in findings[0]["content"]
    finally:
        os.unlink(path)


# ---------------------------------------------------------------------------
# parse_docx
# ---------------------------------------------------------------------------


def _write_docx(sections: list) -> str:
    """Create a temp docx. sections is [(heading_or_None, [paragraphs])].

    heading_or_None=None means no heading paragraph before the body paragraphs.
    heading_or_None=str adds a Heading 1 paragraph with that text first.
    """
    import docx as python_docx

    document = python_docx.Document()
    for heading, paragraphs in sections:
        if heading is not None:
            document.add_heading(heading, level=1)
        for p in paragraphs:
            document.add_paragraph(p)
    fh = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
    fh.close()
    document.save(fh.name)
    return fh.name


def test_parse_docx_headings_become_titles():
    path = _write_docx([("Introduction", ["Some intro text."]), ("Conclusion", ["Wrap up."])])
    try:
        findings = parse_docx(path, "essay.docx")
        titles = [f["title"] for f in findings]
        assert "Introduction" in titles
        assert "Conclusion" in titles
    finally:
        os.unlink(path)


def test_parse_docx_content_prefix_has_section():
    path = _write_docx([("Background", ["Background content here."])])
    try:
        findings = parse_docx(path, "doc.docx")
        bg = next(f for f in findings if f["title"] == "Background")
        assert "[File: doc.docx | Section: Background]" in bg["content"]
    finally:
        os.unlink(path)


def test_parse_docx_all_findings_have_file_upload_source():
    path = _write_docx([("Section A", ["Text A."]), ("Section B", ["Text B."])])
    try:
        findings = parse_docx(path, "file.docx")
        assert all(f["source"] == "file_upload" for f in findings)
    finally:
        os.unlink(path)


def test_parse_docx_no_heading_falls_back_to_section_n():
    path = _write_docx([(None, ["Just some text without a heading."])])
    try:
        findings = parse_docx(path, "plain.docx")
        # Should fall back to "plain.docx - Section 1"
        assert any("Section 1" in f["title"] for f in findings)
    finally:
        os.unlink(path)


def test_parse_docx_empty_sections_skipped():
    path = _write_docx([("EmptySection", []), ("RealSection", ["Actual content."])])
    try:
        findings = parse_docx(path, "mixed.docx")
        titles = [f["title"] for f in findings]
        assert "EmptySection" not in titles
        assert "RealSection" in titles
    finally:
        os.unlink(path)


def test_parse_docx_paragraph_text_in_content():
    path = _write_docx([("Details", ["Marker sentence alpha beta gamma."])])
    try:
        findings = parse_docx(path, "detail.docx")
        detail = next(f for f in findings if f["title"] == "Details")
        assert "Marker sentence alpha beta gamma." in detail["content"]
    finally:
        os.unlink(path)


# ---------------------------------------------------------------------------
# parse_file dispatcher
# ---------------------------------------------------------------------------


def test_parse_file_routes_csv():
    path = _write_csv([{"col": "val"}])
    try:
        findings = parse_file(path, "data.csv")
        assert len(findings) >= 1
        assert all(f["source"] == "file_upload" for f in findings)
    finally:
        os.unlink(path)


def test_parse_file_routes_xlsx():
    path = _write_xlsx({"Sheet1": [{"a": 1}]})
    try:
        findings = parse_file(path, "data.xlsx")
        assert len(findings) >= 1
    finally:
        os.unlink(path)


def test_parse_file_routes_xls_same_as_xlsx():
    path = _write_xlsx({"Sheet1": [{"a": 1}]})
    xls_path = path.replace(".xlsx", ".xls")
    os.rename(path, xls_path)
    try:
        findings = parse_file(xls_path, "data.xls")
        assert len(findings) >= 1
    finally:
        os.unlink(xls_path)


def test_parse_file_routes_pdf():
    path = _write_pdf(["Hello PDF"])
    try:
        findings = parse_file(path, "doc.pdf")
        assert len(findings) >= 1
    finally:
        os.unlink(path)


def test_parse_file_routes_docx():
    path = _write_docx([("Title", ["Body text."])])
    try:
        findings = parse_file(path, "file.docx")
        assert len(findings) >= 1
    finally:
        os.unlink(path)


def test_parse_file_routes_txt():
    fh = tempfile.NamedTemporaryFile(mode="w", suffix=".txt", delete=False)
    fh.write("Some plain text content here.")
    fh.close()
    try:
        findings = parse_file(fh.name, "notes.txt")
        assert len(findings) >= 1
    finally:
        os.unlink(fh.name)


def test_parse_file_raises_for_unsupported_extension():
    fh = tempfile.NamedTemporaryFile(suffix=".zip", delete=False)
    fh.close()
    try:
        with pytest.raises(ValueError, match=r"\.zip"):
            parse_file(fh.name, "archive.zip")
    finally:
        os.unlink(fh.name)
