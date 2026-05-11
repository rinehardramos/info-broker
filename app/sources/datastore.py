"""Data-plane storage for uploaded research sources.

System Postgres stores source metadata and artifact pointers only. User/client
data is normalized into Parquet artifacts under SOURCE_DATASTORE_ROOT and queried
through DuckDB so large tabular files do not have to be loaded into API memory.
"""

from __future__ import annotations

import os
import re
import uuid
from pathlib import Path
from typing import Iterable

import pyarrow as pa
import pyarrow.parquet as pq

_TABULAR_SUFFIXES = {".csv", ".xlsx", ".xls", ".parquet"}
_BATCH_SIZE = 1000
_QUERY_CANDIDATE_LIMIT = 5000
_ROW_RESULT_LIMIT = 50
_META_COLUMNS = {"__row_number", "__sheet_name", "__search_text"}

_ROW_PRIORITY_PATTERNS = [
    "fullname", "firstname", "lastname", "headline", "publicidentifier", "linkedinurl",
    "currentposition/0/companyname", "currentposition/0/companylinkedinurl",
    "currentposition/0/daterange/start/year", "companywebsites/0/domain",
    "companywebsites/0/url", "companywebsites/0/validemailserver",
    "companywebsites/1/domain", "companywebsites/1/validemailserver",
    "location/parsed/country", "location/parsed/countryfull", "location/linkedintext",
    "emails/0/email", "emails/0/deliverable", "emails/0/status",
]
_SME_OUTSOURCE_TERMS = {"sme", "small", "medium", "outsourced", "outsourcing", "third", "party", "it", "msp"}
_DECISION_MAKER_TERMS = [
    "founder", "co-founder", "owner", "ceo", "president", "managing director",
    "partner", "principal", "director", "operations", "office manager",
    "head of it", "it manager", "technology",
]
_TECH_COMPETITOR_TERMS = [
    "software", "saas", "it services", "managed services", "msp", "cybersecurity",
    "cloud", "systems integrator", "app development", "web development",
]


def is_tabular_file(filename: str) -> bool:
    return Path(filename).suffix.lower() in _TABULAR_SUFFIXES


def datastore_root() -> Path:
    return Path(os.getenv("SOURCE_DATASTORE_ROOT", "/tmp/info-broker-datastore")).expanduser()


def ingest_tabular_source(source_id: str, user_id: str, file_path: str, filename: str) -> dict:
    """Normalize a tabular upload into data-plane Parquet artifacts.

    Returns a manifest that contains table metadata and artifact pointers, but no
    row values. That keeps user data out of the control-plane database.
    """
    source_dir = datastore_root() / str(user_id) / str(source_id)
    source_dir.mkdir(parents=True, exist_ok=True)

    ext = Path(filename).suffix.lower()
    if ext == ".csv":
        tables = [_ingest_csv(source_dir, file_path)]
    elif ext == ".xlsx":
        tables = _ingest_xlsx(source_dir, file_path)
    elif ext == ".xls":
        tables = _ingest_xls(source_dir, file_path)
    elif ext == ".parquet":
        tables = [_ingest_parquet(source_dir, file_path)]
    else:
        raise ValueError(f"Unsupported tabular file extension: {ext!r}")

    return {
        "storage": "data_plane",
        "artifact_format": "parquet",
        "datastore_root": str(datastore_root()),
        "source_artifact_dir": str(source_dir),
        "filename": filename,
        "tables": tables,
        "table_count": len(tables),
        "row_count": sum(int(t["row_count"]) for t in tables),
    }


def query_tabular_source(source: dict, query: str, limit: int = 20) -> list[dict]:
    manifest = source.get("manifest") or {}
    if isinstance(manifest, str):
        import json
        manifest = json.loads(manifest)
    if manifest.get("storage") != "data_plane":
        return []

    query_terms = _tokenize(query)
    if not query_terms:
        return []

    import duckdb

    regex = "|".join(re.escape(term) for term in sorted(query_terms, key=len, reverse=True))
    candidate_limit = min(max(limit * 100, 500), _QUERY_CANDIDATE_LIMIT)
    ranked: list[tuple[float, dict]] = []

    with duckdb.connect(database=":memory:") as conn:
        for table in manifest.get("tables", []):
            parquet_path = table.get("artifact_path", "")
            if not parquet_path or not Path(parquet_path).exists():
                continue

            rows = conn.execute(
                """
                SELECT *
                  FROM read_parquet(?)
                 WHERE regexp_matches(lower(__search_text), ?)
                 LIMIT ?
                """,
                [parquet_path, regex, candidate_limit],
            ).fetchall()
            columns = [desc[0] for desc in conn.description or []]
            for values in rows:
                raw = dict(zip(columns, values))
                row_data = _select_response_columns(raw, query_terms)
                if not row_data:
                    continue
                score = _score_row(row_data, query_terms)
                if score <= 0:
                    continue
                row_number = int(raw.get("__row_number") or 0)
                sheet = str(raw.get("__sheet_name") or table.get("sheet_name") or "Data")
                ranked.append((
                    score,
                    _format_result(
                        source=source,
                        sheet=sheet,
                        row_number=row_number,
                        row_data=row_data,
                        score=score,
                    ),
                ))

    ranked.sort(key=lambda item: item[0], reverse=True)
    return [item for _score, item in ranked[: min(limit, _ROW_RESULT_LIMIT)]]


def tabular_manifest_findings(filename: str, manifest: dict) -> list[dict]:
    """Create schema-only findings for semantic discovery in Qdrant."""
    findings: list[dict] = []
    for table in manifest.get("tables", []):
        sheet = table.get("sheet_name", "Data")
        columns = table.get("columns", [])
        column_lines = "\n".join(f"- **{col}**" for col in columns)
        content = (
            f"[File: {filename} | {sheet} Schema]\n"
            f"**Storage:** data-plane parquet\n"
            f"**Rows:** {table.get('row_count', 0)}\n"
            f"**Columns ({len(columns)}):**\n{column_lines}"
        )
        findings.append({
            "title": f"{filename} - {sheet} Schema",
            "content": content,
            "source": "file_upload",
            "confidence": 95,
        })
    return findings


def _ingest_csv(source_dir: Path, file_path: str) -> dict:
    import pandas as pd

    writer: pq.ParquetWriter | None = None
    columns: list[str] = []
    row_count = 0
    artifact_path = source_dir / "data.parquet"
    try:
        for chunk in pd.read_csv(file_path, chunksize=_BATCH_SIZE, dtype=str, keep_default_na=False):
            if not columns:
                columns = _normalize_columns(chunk.columns)
                chunk.columns = columns
            else:
                chunk.columns = columns
            records = chunk.to_dict(orient="records")
            table = _records_to_arrow(records, columns, "Data", row_count + 1)
            writer = _write_table(writer, artifact_path, table)
            row_count += len(records)
    finally:
        if writer:
            writer.close()

    if not columns:
        columns = []
        _write_empty_table(artifact_path, columns)
    return _table_manifest("Data", artifact_path, columns, row_count)


def _ingest_xlsx(source_dir: Path, file_path: str) -> list[dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    tables: list[dict] = []
    try:
        for sheet in workbook.worksheets:
            tables.append(_ingest_row_iter(source_dir, sheet.title, sheet.iter_rows(values_only=True)))
    finally:
        workbook.close()
    return tables


def _ingest_xls(source_dir: Path, file_path: str) -> list[dict]:
    import xlrd

    book = xlrd.open_workbook(file_path, on_demand=True)
    tables: list[dict] = []
    try:
        for sheet_name in book.sheet_names():
            sheet = book.sheet_by_name(sheet_name)
            rows = (sheet.row_values(i) for i in range(sheet.nrows))
            tables.append(_ingest_row_iter(source_dir, sheet_name, rows))
    finally:
        book.release_resources()
    return tables


def _ingest_parquet(source_dir: Path, file_path: str) -> dict:
    artifact_path = source_dir / "data.parquet"
    parquet = pq.ParquetFile(file_path)
    source_columns = list(parquet.schema.names)
    columns = _normalize_columns(source_columns)
    writer: pq.ParquetWriter | None = None
    row_count = 0
    try:
        for batch in parquet.iter_batches(batch_size=_BATCH_SIZE):
            raw = batch.to_pydict()
            records = []
            for idx in range(batch.num_rows):
                records.append({
                    dst: _clean_cell(raw[src][idx])
                    for src, dst in zip(source_columns, columns)
                })
            table = _records_to_arrow(records, columns, "Data", row_count + 1)
            writer = _write_table(writer, artifact_path, table)
            row_count += len(records)
    finally:
        if writer:
            writer.close()
    if row_count == 0:
        _write_empty_table(artifact_path, columns)
    return _table_manifest("Data", artifact_path, columns, int(row_count))


def _ingest_row_iter(source_dir: Path, sheet_name: str, rows: Iterable[Iterable[object]]) -> dict:
    safe_sheet = re.sub(r"[^A-Za-z0-9_.-]+", "_", sheet_name).strip("_") or "sheet"
    artifact_path = source_dir / f"{safe_sheet}-{uuid.uuid4().hex[:8]}.parquet"
    iterator = iter(rows)
    try:
        header = next(iterator)
    except StopIteration:
        _write_empty_table(artifact_path, [])
        return _table_manifest(sheet_name, artifact_path, [], 0)

    columns = _normalize_columns(header)
    batch: list[dict[str, str]] = []
    writer: pq.ParquetWriter | None = None
    row_count = 0
    try:
        for values in iterator:
            record = {
                col: _clean_cell(values[idx] if idx < len(values) else "")
                for idx, col in enumerate(columns)
            }
            batch.append(record)
            if len(batch) >= _BATCH_SIZE:
                table = _records_to_arrow(batch, columns, sheet_name, row_count + 1)
                writer = _write_table(writer, artifact_path, table)
                row_count += len(batch)
                batch = []
        if batch:
            table = _records_to_arrow(batch, columns, sheet_name, row_count + 1)
            writer = _write_table(writer, artifact_path, table)
            row_count += len(batch)
    finally:
        if writer:
            writer.close()

    if row_count == 0:
        _write_empty_table(artifact_path, columns)
    return _table_manifest(sheet_name, artifact_path, columns, row_count)


def _records_to_arrow(records: list[dict[str, str]], columns: list[str], sheet_name: str, start_row: int) -> pa.Table:
    data: dict[str, list[str | int]] = {
        "__row_number": list(range(start_row, start_row + len(records))),
        "__sheet_name": [sheet_name] * len(records),
    }
    for col in columns:
        data[col] = [record.get(col, "") for record in records]
    data["__search_text"] = [
        " ".join(str(record.get(col, "")).lower() for col in columns)
        for record in records
    ]
    schema = pa.schema(
        [pa.field("__row_number", pa.int64()), pa.field("__sheet_name", pa.string())]
        + [pa.field(col, pa.string()) for col in columns]
        + [pa.field("__search_text", pa.string())]
    )
    return pa.Table.from_pydict(data, schema=schema)


def _write_table(writer: pq.ParquetWriter | None, path: Path, table: pa.Table) -> pq.ParquetWriter:
    if writer is None:
        writer = pq.ParquetWriter(path, table.schema, compression="zstd")
    writer.write_table(table)
    return writer


def _write_empty_table(path: Path, columns: list[str]) -> None:
    schema = pa.schema(
        [pa.field("__row_number", pa.int64()), pa.field("__sheet_name", pa.string())]
        + [pa.field(col, pa.string()) for col in columns]
        + [pa.field("__search_text", pa.string())]
    )
    pq.write_table(pa.Table.from_batches([], schema=schema), path, compression="zstd")


def _table_manifest(sheet_name: str, artifact_path: Path, columns: list[str], row_count: int) -> dict:
    return {
        "table_id": str(uuid.uuid4()),
        "sheet_name": sheet_name,
        "row_count": row_count,
        "columns": columns,
        "artifact_path": str(artifact_path),
    }


def _normalize_columns(columns: Iterable[object]) -> list[str]:
    seen: dict[str, int] = {}
    normalized: list[str] = []
    for idx, column in enumerate(columns, start=1):
        base = _clean_cell(column) or f"column_{idx}"
        if base in _META_COLUMNS:
            base = f"source_{base.strip('_')}"
        count = seen.get(base, 0) + 1
        seen[base] = count
        normalized.append(base if count == 1 else f"{base}__{count}")
    return normalized


def _clean_cell(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() in {"nan", "none", "nat"}:
        return ""
    return text.replace("\x00", "").replace("\n", " ")[:1000]


def _tokenize(text: str) -> set[str]:
    return {t for t in re.split(r"[^a-z0-9]+", text.lower()) if len(t) >= 2}


def _select_response_columns(row: dict, query_terms: set[str]) -> dict[str, str]:
    selected: dict[str, str] = {}
    fallback: dict[str, str] = {}
    for key, value in row.items():
        if key in _META_COLUMNS or value in (None, ""):
            continue
        text = _clean_cell(value)
        if not text:
            continue
        key_lower = key.lower()
        if key_lower.startswith("moreprofiles/"):
            continue
        compact = key_lower.replace("_", "").replace(" ", "")
        if any(pattern in compact for pattern in _ROW_PRIORITY_PATTERNS) or query_terms.intersection(_tokenize(key_lower)):
            selected[key] = text
        elif len(fallback) < 12:
            fallback[key] = text
    if not selected:
        selected = fallback
    return dict(list(selected.items())[:80])


def _score_row(row_data: dict[str, str], query_terms: set[str]) -> float:
    haystack = " ".join(f"{k} {v}" for k, v in row_data.items()).lower()
    lexical_score = sum(1 for term in query_terms if term in haystack)
    heuristic_score = _sme_outsource_row_score(row_data) if _SME_OUTSOURCE_TERMS.intersection(query_terms) else 0
    return float(lexical_score + heuristic_score)


def _sme_outsource_row_score(row_data: dict[str, str]) -> float:
    text = " ".join(row_data.values()).lower()
    score = 0.0
    if any(term in text for term in _DECISION_MAKER_TERMS):
        score += 6
    if any(term in text for term in ("founder", "owner", "ceo", "president", "managing director")):
        score += 5
    if "true" in str(row_data.get("companyWebsites/0/validEmailServer", "")).lower():
        score += 2
    if any(term in text for term in _TECH_COMPETITOR_TERMS):
        score -= 4
    return score


def _format_result(source: dict, sheet: str, row_number: int, row_data: dict[str, str], score: float) -> dict:
    filename = source.get("filename", "upload")
    name = " ".join(v for v in [row_data.get("firstName", ""), row_data.get("lastName", "")] if v).strip()
    company = row_data.get("currentPosition/0/companyName", "")
    label = " / ".join(part for part in [name, company] if part) or f"row {row_number}"
    lines = [f"[File: {filename} | Sheet: {sheet} | Row: {row_number}]"]
    lines.extend(f"- {key}: {value}" for key, value in row_data.items())
    return {
        "title": f"{filename} - {sheet} row {row_number}: {label}",
        "content": "\n".join(lines)[:4000],
        "score": round(score, 4),
        "source_tool": "file_upload_row",
        "run_id": str(source.get("id", "")),
        "row_number": row_number,
        "sheet": sheet,
        "filename": filename,
        "row_data": row_data,
    }
