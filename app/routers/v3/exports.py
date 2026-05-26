"""Export API — trigger and download research exports (PDF, CSV, Excel)."""

from __future__ import annotations

import json
import logging
import os
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse
from pydantic import BaseModel

from app.routers.v3.auth import get_current_user
from app.routers.v3.tenancy import org_scope_clause
from app.routers.v3.db import fetch_one

router = APIRouter(prefix="/v3/exports", tags=["v3-exports"])
log = logging.getLogger(__name__)

_EXPORTS_DIR = "/tmp/exports"

_MEDIA_TYPES = {
    "pdf": "application/pdf",
    "csv": "text/csv",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "json": "application/json",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class ExportRequest(BaseModel):
    format: str = "pdf"  # "pdf" | "csv" | "xlsx" | "json" | "docx"
    include_analysis: bool = True


def _normalize_finding(f: dict) -> dict:
    """Map a research_trails finding (candidate_name / source_url / source_class /
    evidence_snippet / phase_id …) onto the export shape. The generators used to
    read title/source/content/url, which the real findings don't have — hence the
    near-empty exports. This bridges both shapes."""
    return {
        "title": f.get("title") or f.get("candidate_name") or f.get("candidate") or f.get("name") or "",
        "source_class": f.get("source") or f.get("source_class") or "",
        "url": f.get("url") or f.get("source_url") or "",
        "confidence": f.get("confidence", ""),
        "phase": f.get("phase_id") or f.get("phase") or "",
        "content": (
            f.get("content") or f.get("evidence_snippet") or f.get("evidence_summary")
            or f.get("claim") or ""
        ),
    }


@router.post("/research/{run_id}")
def trigger_export(
    run_id: str,
    body: ExportRequest,
    user: dict = Depends(get_current_user),
) -> dict:
    """Generate an export file for a research trail run and return its download URL."""
    fmt = body.format.lower()
    if fmt not in _MEDIA_TYPES:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported format: {fmt!r}. Use pdf, csv, xlsx, json, or docx.",
        )

    # Fetch research trail + run metadata by run_id
    _clause, _cparams = org_scope_clause(user)
    row = fetch_one(
        f"SELECT rt.query, rt.findings, rt.analysis, rt.trail, rt.entity_type, rt.tool_calls, "  # noqa: S608 - clause is a constant org-scope fragment; values parameterized
        f"pr.status, pr.started_at, pr.finished_at "
        f"FROM research_trails rt "
        f"JOIN pipeline_runs pr ON pr.id = rt.run_id "
        f"WHERE rt.run_id = %s {_clause.replace('AND org_id', 'AND pr.org_id')}",
        tuple([run_id, *_cparams]),
    )
    if not row:
        raise HTTPException(status_code=404, detail=f"Research trail not found for run_id={run_id!r}")

    def _loads(v):
        return json.loads(v) if isinstance(v, str) else v

    findings = [_normalize_finding(f) for f in (_loads(row.get("findings")) or [])]
    analysis = _loads(row.get("analysis")) if body.include_analysis else None
    trail = _loads(row.get("trail")) or {}
    ranked = trail.get("ranked_candidates") or [] if isinstance(trail, dict) else []
    meta = {
        "run_id": run_id,
        "query": row.get("query", "Research Report"),
        "entity_type": row.get("entity_type"),
        "status": row.get("status"),
        "started_at": str(row.get("started_at")) if row.get("started_at") else None,
        "finished_at": str(row.get("finished_at")) if row.get("finished_at") else None,
        "tool_calls": row.get("tool_calls"),
        "phases": trail.get("phases") if isinstance(trail, dict) else None,
        "finding_count": len(findings),
        "candidate_count": len(ranked),
    }
    query = meta["query"]

    os.makedirs(_EXPORTS_DIR, exist_ok=True)
    output_path = f"{_EXPORTS_DIR}/{run_id}.{fmt}"

    try:
        if fmt == "pdf":
            _generate_pdf(meta, findings, ranked, analysis, output_path)
        elif fmt == "csv":
            _generate_csv(findings, output_path)
        elif fmt == "xlsx":
            _generate_excel(findings, ranked, analysis, meta, output_path)
        elif fmt == "json":
            _generate_json(meta, findings, ranked, analysis, output_path)
        elif fmt == "docx":
            _generate_docx(meta, findings, ranked, analysis, output_path)
    except Exception as exc:
        log.error("Export generation failed for run_id=%s fmt=%s: %s", run_id, fmt, exc)
        raise HTTPException(status_code=500, detail=f"Export generation failed: {exc}")

    filename = f"{run_id}.{fmt}"
    return {"filename": filename, "url": f"/v3/exports/files/{filename}"}


@router.get("/files/{filename}")
def download_export(filename: str) -> FileResponse:
    """Serve an export file. No auth required — files are ephemeral and URL is the token."""
    # Prevent path traversal
    safe_name = Path(filename).name
    if safe_name != filename or ".." in filename:
        raise HTTPException(status_code=400, detail="Invalid filename")

    file_path = f"{_EXPORTS_DIR}/{safe_name}"
    if not os.path.exists(file_path):
        raise HTTPException(status_code=404, detail="Export file not found")

    ext = safe_name.rsplit(".", 1)[-1].lower() if "." in safe_name else ""
    media_type = _MEDIA_TYPES.get(ext, "application/octet-stream")

    return FileResponse(path=file_path, media_type=media_type, filename=safe_name)


# ---------------------------------------------------------------------------
# Internal generators (same logic as the pipeline nodes, but callable directly)
# ---------------------------------------------------------------------------

def _generate_pdf(meta: dict, findings: list[dict], ranked: list[dict], analysis: dict | None, output_path: str) -> None:
    from fpdf import FPDF
    from datetime import datetime

    date_str = datetime.now().strftime("%Y-%m-%d %H:%M")

    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_left_margin(10)
    pdf.set_right_margin(10)
    pdf.add_page()

    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "Research Report", ln=True)
    pdf.set_font("Helvetica", "", 11)
    _mc(pdf, 7, _safe(f"Query: {meta.get('query', '')}"))
    pdf.set_font("Helvetica", "", 9)
    pdf.set_text_color(110, 110, 110)
    for label, key in (("Run ID", "run_id"), ("Status", "status"), ("Entity type", "entity_type"),
                       ("Started", "started_at"), ("Finished", "finished_at")):
        if meta.get(key):
            pdf.cell(0, 5, _safe(f"{label}: {meta[key]}"), ln=True)
    phases = meta.get("phases")
    if phases:
        pdf.cell(0, 5, _safe(f"Phases: {' -> '.join(str(p) for p in phases)}"), ln=True)
    pdf.cell(0, 5, _safe(f"Findings: {meta.get('finding_count', len(findings))} | "
                         f"Ranked candidates: {meta.get('candidate_count', len(ranked))} | "
                         f"Tool calls: {meta.get('tool_calls', '?')} | Generated: {date_str}"), ln=True)
    pdf.set_text_color(0, 0, 0)
    pdf.ln(4)

    # Ranked candidates (the synthesized leads/answers)
    if ranked:
        pdf.set_font("Helvetica", "B", 13)
        pdf.cell(0, 9, "RANKED CANDIDATES", ln=True)
        pdf.line(pdf.get_x(), pdf.get_y(), pdf.get_x() + 190, pdf.get_y())
        pdf.ln(3)
        for i, c in enumerate(ranked, 1):
            pdf.set_font("Helvetica", "B", 10)
            _mc(pdf, 6, _safe(f"{i}. {c.get('name', '')}  (conf {c.get('confidence', '?')})"))
            ev = c.get("evidence") or []
            pdf.set_font("Helvetica", "", 9)
            for e in (ev if isinstance(ev, list) else [])[:3]:
                if isinstance(e, dict) and e.get("snippet"):
                    pdf.set_x(pdf.l_margin)
                    _mc(pdf, 5, _safe("   - " + _trunc(e["snippet"], 220)))
                    if e.get("source_url"):
                        pdf.set_font("Helvetica", "I", 8)
                        pdf.set_text_color(60, 100, 180)
                        _mc(pdf, 5, _safe("     " + _trunc(e["source_url"], 110)))
                        pdf.set_text_color(0, 0, 0)
                        pdf.set_font("Helvetica", "", 9)
            pdf.ln(1)
        pdf.ln(2)

    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 9, "FINDINGS", ln=True)
    pdf.line(pdf.get_x(), pdf.get_y(), pdf.get_x() + 190, pdf.get_y())
    pdf.ln(3)

    if not findings:
        pdf.set_font("Helvetica", "", 10)
        pdf.cell(0, 7, "No findings.", ln=True)
    else:
        for i, f in enumerate(findings, 1):
            pdf.set_font("Helvetica", "B", 10)
            _mc(pdf, 6, _safe(f"{i}. {f.get('title') or f'Finding {i}'}"))
            pdf.set_font("Helvetica", "", 9)
            parts = []
            if f.get("source_class"):
                parts.append(f"Source: {f['source_class']}")
            if f.get("phase"):
                parts.append(f"Phase: {f['phase']}")
            if f.get("confidence") not in (None, ""):
                parts.append(f"Confidence: {f['confidence']}")
            if parts:
                pdf.set_text_color(120, 120, 120)
                pdf.cell(0, 6, _safe("   " + " | ".join(parts)), ln=True)
                pdf.set_text_color(0, 0, 0)
            if f.get("content"):
                pdf.set_x(pdf.l_margin)
                _mc(pdf, 5, _safe("   " + _trunc(f["content"], 400)))
            if f.get("url"):
                pdf.set_font("Helvetica", "I", 8)
                pdf.set_text_color(60, 100, 180)
                _mc(pdf, 5, _safe("   " + _trunc(f["url"], 110)))
                pdf.set_text_color(0, 0, 0)
            pdf.ln(2)

    if analysis:
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 13)
        pdf.cell(0, 9, "ANALYSIS", ln=True)
        pdf.line(pdf.get_x(), pdf.get_y(), pdf.get_x() + 190, pdf.get_y())
        pdf.ln(3)

        for ent in analysis.get("entities", []):
            pdf.set_font("Helvetica", "B", 11)
            pdf.cell(0, 8, "Entities", ln=True)
            pdf.set_font("Helvetica", "", 9)
            name = ent.get("name", "")
            etype = ent.get("type", "")
            pdf.cell(0, 6, _safe(f"  - {name}" + (f" ({etype})" if etype else "")), ln=True)
            break

        for ent in analysis.get("entities", [])[1:]:
            name = ent.get("name", "")
            etype = ent.get("type", "")
            pdf.set_font("Helvetica", "", 9)
            pdf.cell(0, 6, _safe(f"  - {name}" + (f" ({etype})" if etype else "")), ln=True)

        if analysis.get("relationships"):
            pdf.ln(2)
            pdf.set_font("Helvetica", "B", 11)
            pdf.cell(0, 8, "Relationships", ln=True)
            pdf.set_font("Helvetica", "", 9)
            for rel in analysis["relationships"]:
                frm = rel.get("from", rel.get("from_entity", ""))
                to = rel.get("to", rel.get("to_entity", ""))
                rtype = rel.get("type", rel.get("relationship_type", ""))
                pdf.cell(0, 6, _safe(f"  - {frm} --[{rtype}]--> {to}"), ln=True)

        insights = analysis.get("insights", analysis.get("key_insights", []))
        if insights:
            pdf.ln(2)
            pdf.set_font("Helvetica", "B", 11)
            pdf.cell(0, 8, "Insights", ln=True)
            pdf.set_font("Helvetica", "", 9)
            for ins in insights:
                text = ins if isinstance(ins, str) else ins.get("text", ins.get("insight", str(ins)))
                pdf.set_x(pdf.l_margin)
                _mc(pdf, 5, f"  - {_trunc(text, 300)}")

        recs = analysis.get("recommendations", [])
        if recs:
            pdf.ln(2)
            pdf.set_font("Helvetica", "B", 11)
            pdf.cell(0, 8, "Recommendations", ln=True)
            pdf.set_font("Helvetica", "", 9)
            for rec in recs:
                text = rec if isinstance(rec, str) else rec.get("text", rec.get("recommendation", str(rec)))
                pdf.set_x(pdf.l_margin)
                _mc(pdf, 5, f"  - {_trunc(text, 300)}")

    pdf.output(output_path)


_FINDING_COLS = ["rank", "title", "source_class", "url", "confidence", "phase", "content"]


def _finding_rows(findings: list[dict]) -> list[dict]:
    return [
        {
            "rank": i,
            "title": f.get("title", ""),
            "source_class": f.get("source_class", ""),
            "url": f.get("url", ""),
            "confidence": f.get("confidence", ""),
            "phase": f.get("phase", ""),
            "content": f.get("content", ""),
        }
        for i, f in enumerate(findings, 1)
    ]


def _generate_csv(findings: list[dict], output_path: str) -> None:
    import pandas as pd

    df = pd.DataFrame(_finding_rows(findings), columns=_FINDING_COLS)
    df.to_csv(output_path, index=False)


def _generate_excel(findings: list[dict], ranked: list[dict], analysis: dict | None,
                    meta: dict, output_path: str) -> None:
    import pandas as pd

    df_findings = pd.DataFrame(_finding_rows(findings), columns=_FINDING_COLS)

    entities = analysis.get("entities", []) if analysis else []
    entity_rows = [
        {
            "name": e.get("name", ""),
            "type": e.get("type", e.get("entity_type", "")),
            "attributes": str(e.get("attributes", e.get("properties", ""))),
        }
        for e in entities
    ]
    df_entities = pd.DataFrame(entity_rows, columns=["name", "type", "attributes"])

    relationships = analysis.get("relationships", []) if analysis else []
    rel_rows = [
        {
            "from": r.get("from", r.get("from_entity", "")),
            "to": r.get("to", r.get("to_entity", "")),
            "type": r.get("type", r.get("relationship_type", "")),
            "evidence": r.get("evidence", r.get("description", "")),
        }
        for r in relationships
    ]
    df_relationships = pd.DataFrame(rel_rows, columns=["from", "to", "type", "evidence"])

    insights = analysis.get("insights", analysis.get("key_insights", [])) if analysis else []
    recs = analysis.get("recommendations", []) if analysis else []
    insight_rows = [
        {"type": "insight", "text": ins if isinstance(ins, str) else ins.get("text", ins.get("insight", str(ins)))}
        for ins in insights
    ] + [
        {"type": "recommendation", "text": rec if isinstance(rec, str) else rec.get("text", rec.get("recommendation", str(rec)))}
        for rec in recs
    ]
    df_insights = pd.DataFrame(insight_rows, columns=["type", "text"])

    df_meta = pd.DataFrame([{"field": k, "value": str(v)} for k, v in meta.items()],
                           columns=["field", "value"])
    ranked_rows = [
        {
            "rank": i,
            "name": c.get("name", ""),
            "confidence": c.get("confidence", ""),
            "evidence": " | ".join(
                e.get("snippet", "") for e in (c.get("evidence") or []) if isinstance(e, dict)
            )[:2000],
            "sources": " ; ".join(
                e.get("source_url", "") for e in (c.get("evidence") or [])
                if isinstance(e, dict) and e.get("source_url")
            ),
        }
        for i, c in enumerate(ranked, 1)
    ]
    df_ranked = pd.DataFrame(ranked_rows, columns=["rank", "name", "confidence", "evidence", "sources"])

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df_meta.to_excel(writer, sheet_name="Run Info", index=False)
        df_ranked.to_excel(writer, sheet_name="Ranked Candidates", index=False)
        df_findings.to_excel(writer, sheet_name="Findings", index=False)
        df_entities.to_excel(writer, sheet_name="Entities", index=False)
        df_relationships.to_excel(writer, sheet_name="Relationships", index=False)
        df_insights.to_excel(writer, sheet_name="Insights", index=False)


def _generate_json(meta: dict, findings: list[dict], ranked: list[dict],
                   analysis: dict | None, output_path: str) -> None:
    """Full structured export — everything the run produced, machine-readable."""
    report = {
        "metadata": meta,
        "findings": findings,
        "ranked_candidates": ranked,
        "analysis": analysis,
    }
    with open(output_path, "w", encoding="utf-8") as fh:
        json.dump(report, fh, indent=2, ensure_ascii=False, default=str)


def _generate_docx(meta: dict, findings: list[dict], ranked: list[dict],
                   analysis: dict | None, output_path: str) -> None:
    """Formatted Word report (.docx) — headings, run info, ranked candidates,
    findings, and analysis."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading("Research Report", level=0)
    doc.add_paragraph(str(meta.get("query", "")))

    info = doc.add_paragraph()
    info.add_run("Run details\n").bold = True
    for label, key in (("Run ID", "run_id"), ("Status", "status"), ("Entity type", "entity_type"),
                       ("Started", "started_at"), ("Finished", "finished_at"),
                       ("Tool calls", "tool_calls"), ("Findings", "finding_count"),
                       ("Ranked candidates", "candidate_count")):
        if meta.get(key) not in (None, ""):
            r = info.add_run(f"{label}: {meta[key]}\n")
            r.font.size = Pt(9)
    if meta.get("phases"):
        info.add_run(f"Phases: {' -> '.join(str(p) for p in meta['phases'])}\n").font.size = Pt(9)

    if ranked:
        doc.add_heading("Ranked Candidates", level=1)
        for i, c in enumerate(ranked, 1):
            p = doc.add_paragraph(style="List Number")
            p.add_run(f"{c.get('name', '')}  ").bold = True
            p.add_run(f"(confidence {c.get('confidence', '?')})").italic = True
            for e in (c.get("evidence") or [])[:3]:
                if isinstance(e, dict) and e.get("snippet"):
                    sub = doc.add_paragraph(e["snippet"], style="List Bullet 2")
                    if e.get("source_url"):
                        sub.add_run(f"  [{e['source_url']}]").italic = True

    doc.add_heading("Findings", level=1)
    if not findings:
        doc.add_paragraph("No findings.")
    for i, f in enumerate(findings, 1):
        p = doc.add_paragraph(style="List Number")
        p.add_run(f.get("title") or f"Finding {i}").bold = True
        bits = []
        if f.get("source_class"):
            bits.append(f"source: {f['source_class']}")
        if f.get("phase"):
            bits.append(f"phase: {f['phase']}")
        if f.get("confidence") not in (None, ""):
            bits.append(f"confidence: {f['confidence']}")
        if bits:
            meta_run = p.add_run(f"  ({', '.join(bits)})")
            meta_run.italic = True
            meta_run.font.size = Pt(8)
        if f.get("content"):
            doc.add_paragraph(f["content"], style="List Bullet 2")
        if f.get("url"):
            doc.add_paragraph(f["url"], style="List Bullet 2").runs[0].italic = True

    if analysis:
        doc.add_heading("Analysis", level=1)
        for ent in analysis.get("entities", []):
            doc.add_paragraph(
                f"{ent.get('name', '')}" + (f" ({ent.get('type', '')})" if ent.get("type") else ""),
                style="List Bullet",
            )
        for rel in analysis.get("relationships", []):
            frm = rel.get("from", rel.get("from_entity", ""))
            to = rel.get("to", rel.get("to_entity", ""))
            rtype = rel.get("type", rel.get("relationship_type", ""))
            doc.add_paragraph(f"{frm} --[{rtype}]--> {to}", style="List Bullet")
        for ins in analysis.get("insights", analysis.get("key_insights", [])):
            doc.add_paragraph(ins if isinstance(ins, str) else ins.get("text", str(ins)), style="List Bullet")

    doc.save(output_path)


def _mc(pdf, h: float, txt: str) -> None:
    """multi_cell that resets x to the left margin and uses CHAR wrapmode, so long
    unbreakable tokens (e.g. URLs) don't raise 'Not enough horizontal space'."""
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(0, h, txt, wrapmode="CHAR")


def _safe(text: str) -> str:
    """Replace non-latin1 chars for fpdf2 Helvetica."""
    return str(text).encode("latin-1", errors="replace").decode("latin-1")


def _trunc(text: str, max_len: int) -> str:
    text = _safe(str(text))
    return text[:max_len - 3] + "..." if len(text) > max_len else text
